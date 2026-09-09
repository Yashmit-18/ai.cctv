"""Multi-employee daily Excel report generator (Phase 12).

Produces a styled ``.xlsx`` workbook with:

* **Sheet 1 -- Summary**: consolidated per-employee productivity table,
  scored against the expected working window.  ``Unknown`` never appears.
* **Sheet 2 -- Activity Timeline**: raw time-stamped state logs.

All numbers come from :mod:`src.productivity` via :mod:`src.analytics` so
the report, dashboard and daemon share one formula.
"""

import logging
import os
from datetime import date, datetime

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import config
from config import REPORT_OUTPUT_DIR, WORK_SCHEDULE
from src import database as db
from src.analytics import employee_day_metrics
from src.employees import EmployeeStore

logger = logging.getLogger("cctv.reporter")

_S = 3600.0
_M = 60.0

_HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center")

_GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
_YELLOW_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
_RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
_TOTAL_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

_THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

_CENTER = Alignment(horizontal="center", vertical="center")


def _fmt_time(ts: str, day_prefix: str) -> str:
    """Trim 'YYYY-MM-DD ' prefix in the In/Out-Time columns."""
    if not ts:
        return ""
    return ts[len(day_prefix) + 1:] if ts.startswith(day_prefix) else ts


def build_summary(metrics: list[dict], day: date) -> pd.DataFrame:
    """Build the per-employee summary DataFrame from metrics rows."""
    rows = []
    day_prefix = day.isoformat()
    for m in metrics:
        pct = m.get("productive_pct")
        rows.append({
            "Employee ID": m["employee_id"],
            "Name": m.get("name", ""),
            "Department": m.get("department", ""),
            "In Time": _fmt_time(m.get("in_time", ""), day_prefix),
            "Out Time": _fmt_time(m.get("out_time", ""), day_prefix),
            "Expected Hours": round((m["expected_working_seconds"] or 0) / _S, 2),
            "Active Hours": m.get("active_hours", 0.0),
            "Phone (Mins)": m.get("phone_mins", 0.0),
            "Away (Mins)": m.get("away_mins", 0.0),
            "Unobserved (Mins)": round((m.get("unobserved_seconds", 0) or 0) / _M, 1),
            "Productivity (%)": pct if pct is not None else "n/a",
            "Status": m.get("status", ""),
        })
    return pd.DataFrame(rows)


def build_timeline(day_logs: list[tuple]) -> pd.DataFrame:
    """Rows from raw logs (Unknown included for monitoring visibility)."""
    rows = []
    for ts, emp_id, state, dur, source in sorted(day_logs, key=lambda r: r[0]):
        rows.append({
            "Timestamp": ts,
            "Employee ID": emp_id,
            "State": state,
            "Duration (s)": round(dur, 2),
            "Source": source or "",
        })
    return pd.DataFrame(rows)


def _style_header(ws):
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _HEADER_ALIGN
        cell.border = _THIN_BORDER
    ws.freeze_panes = "A2"


def _style_data_rows(ws):
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            cell.alignment = _CENTER
            cell.border = _THIN_BORDER


def _auto_widths(ws, min_width: int = 10, max_width: int = 28):
    for col_cells in ws.columns:
        col_letter = get_column_letter(col_cells[0].column)
        lengths = []
        for cell in col_cells:
            val = str(cell.value) if cell.value is not None else ""
            lengths.append(len(val))
        best = max(lengths) + 3 if lengths else min_width
        ws.column_dimensions[col_letter].width = min(max(best, min_width), max_width)


def _total_row(ws, df: pd.DataFrame):
    """Append a totals row beneath the summary data."""
    if df.empty:
        return
    trow = ws.max_row + 2
    ws.cell(row=trow, column=2, value="TOTAL")
    ws.cell(row=trow, column=2).font = Font(bold=True)
    for col in df.columns:
        idx = list(df.columns).index(col) + 1
        key = col
        if key == "Active Hours":
            _set_total(ws, trow, idx, df, key)
        elif key == "Phone (Mins)":
            _set_total(ws, trow, idx, df, key)
        elif key == "Away (Mins)":
            _set_total(ws, trow, idx, df, key)
        elif key == "Expected Hours":
            _set_total(ws, trow, idx, df, key)
    for c in range(1, ws.max_column + 1):
        cell = ws.cell(row=trow, column=c)
        cell.fill = _TOTAL_FILL
        cell.border = _THIN_BORDER


def _set_total(ws, trow, idx, df, key):
    try:
        vals = [float(v) for v in df[key].tolist() if isinstance(v, (int, float))]
    except (TypeError, ValueError):
        vals = []
    ws.cell(row=trow, column=idx, value=round(sum(vals), 2))


def _highlight_scores(ws, df: pd.DataFrame):
    if df.empty:
        return
    col_name = "Productivity (%)"
    if col_name not in df.columns:
        return
    col_idx = list(df.columns).index(col_name) + 1
    for r in range(2, len(df) + 2):
        cell = ws.cell(row=r, column=col_idx)
        if cell.value in (None, "n/a"):
            cell.fill = _YELLOW_FILL
            continue
        try:
            val = float(cell.value)
        except (TypeError, ValueError):
            continue
        if val >= 80:
            cell.fill = _GREEN_FILL
        elif val >= 50:
            cell.fill = _YELLOW_FILL
        else:
            cell.fill = _RED_FILL


def generate_daily_report(conn, report_date: date) -> str:
    """Generate the styled daily report.  Returns the output path."""
    employees = EmployeeStore(conn, WORK_SCHEDULE)
    metrics = employee_day_metrics(conn, report_date, employees=employees)
    day_logs = db.query_day_intervals(conn, report_date.isoformat())

    summary_df = build_summary(metrics, report_date)
    timeline_df = build_timeline(day_logs)

    os.makedirs(REPORT_OUTPUT_DIR, exist_ok=True)
    out_path = os.path.join(
        REPORT_OUTPUT_DIR,
        f"daily_report_{report_date.strftime('%Y-%m-%d')}.xlsx",
    )

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        summary_df.to_excel(writer, index=False, sheet_name="Summary")
        ws1 = writer.sheets["Summary"]
        _style_header(ws1)
        _style_data_rows(ws1)
        _highlight_scores(ws1, summary_df)
        _total_row(ws1, summary_df)
        _auto_widths(ws1)

        # Footer with generation timestamp + working window note
        frow = ws1.max_row + 2
        ws1.cell(row=frow, column=1, value=f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        ws1.cell(row=frow + 1, column=1,
                 value=f"Working window: {WORK_SCHEDULE.get('start')} - {WORK_SCHEDULE.get('end')} "
                       f"(lunch {WORK_SCHEDULE.get('lunch')})")

        if not timeline_df.empty:
            timeline_df.to_excel(writer, index=False, sheet_name="Activity Timeline")
            ws2 = writer.sheets["Activity Timeline"]
            _style_header(ws2)
            _style_data_rows(ws2)
            _auto_widths(ws2)

    logger.info("Report written: %s (%d summary rows, %d log rows)",
                out_path, len(metrics), len(day_logs))
    return out_path
